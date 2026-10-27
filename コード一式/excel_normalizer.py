"""Module for dynamic scanning, data normalization, and Intermediate Representation (IR) generation for Excel files."""

from dataclasses import asdict, dataclass, field
from datetime import date, datetime
import json
from pathlib import Path
import re
from typing import Any, Dict, List, Optional

import openpyxl
from openpyxl.worksheet.worksheet import Worksheet


@dataclass
class NormalizedRecord:
    """Normalized data model for a single row."""

    row_number: int
    data: Dict[str, Any]


@dataclass
class SheetData:
    """Normalized data structure for an individual sheet."""

    sheet_name: str
    is_visible: bool
    headers: List[str]
    records: List[NormalizedRecord] = field(default_factory=list)
    plain_text_summary: str = ""


@dataclass
class NormalizedWorkbook:
    """Intermediate Representation (IR) for the entire workbook."""

    file_name: str
    processed_at: str
    sheets: List[SheetData] = field(default_factory=list)

    def to_dict(self) -> Dict[str, Any]:
        return asdict(self)

    def to_json(self, indent: int = 2) -> str:
        return json.dumps(self.to_dict(), ensure_ascii=False, indent=indent)


class ExcelNormalizer:
    """Class responsible for loading and normalizing multi-sheet Excel files."""

    def __init__(self, ignore_hidden_sheets: bool = True):
        self.ignore_hidden_sheets = ignore_hidden_sheets

    def normalize_value(self, val: Any) -> Any:
        if val is None:
            return ""

        if isinstance(val, (datetime, date)):
            return val.strftime("%Y-%m-%d")

        if isinstance(val, str):
            cleaned = re.sub(r"[\r\t\x00-\x08\x0b-\x0c\x0e-\x1f]", "", val)
            cleaned = re.sub(r"\n+", "\n", cleaned)
            return cleaned.strip()

        if isinstance(val, (int, float)):
            return val

        return str(val).strip()

    def _is_future_date(self, val: Any, today: date) -> bool:
        if isinstance(val, datetime):
            return val.date() > today
        if isinstance(val, date):
            return val > today

        if isinstance(val, str):
            match = re.search(r"(\d{4})[-/](\d{1,2})[-/](\d{1,2})", val)
            if match:
                try:
                    dt = date(
                        int(match.group(1)),
                        int(match.group(2)),
                        int(match.group(3)),
                    )
                    return dt > today
                except ValueError:
                    return False
        return False

    def _extract_sheet(self, ws: Worksheet) -> Optional[SheetData]:
        is_visible = ws.sheet_state == "visible"
        if self.ignore_hidden_sheets and not is_visible:
            return None

        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            return None

        raw_headers = rows[0]
        headers = [
            str(self.normalize_value(h)) or f"col_{idx+1}"
            for idx, h in enumerate(raw_headers)
        ]

        records: List[NormalizedRecord] = []
        text_lines: List[str] = []
        today = date.today()

        for row_idx, row in enumerate(rows[1:], start=2):
            normalized_row: Dict[str, Any] = {}
            line_items: List[str] = []
            has_valid_data = False
            is_future_row = False

            for col_idx, cell_value in enumerate(row):
                if col_idx < len(headers):
                    col_name = headers[col_idx]
                    norm_val = self.normalize_value(cell_value)

                    if (
                        "日" in col_name or "Date" in col_name.capitalize()
                    ) and self._is_future_date(cell_value, today):
                        is_future_row = True
                        break

                    normalized_row[col_name] = norm_val

                    if norm_val != "":
                        has_valid_data = True
                        line_items.append(f"{col_name}: {norm_val}")

            if is_future_row:
                continue

            if has_valid_data:
                # Extract appeal or learning reference info
                task_type = str(normalized_row.get("業務", ""))
                appeal_info = ""
                if "自己学習" in task_type:
                    ref_text = normalized_row.get(
                        "参照テキスト/サイト", ""
                    ) or normalized_row.get("参照テキスト", "")
                    ref_chapter = normalized_row.get("参照章", "")
                    if ref_text or ref_chapter:
                        appeal_info = (
                            f"[Appeal/Learning Ref] Text: {ref_text} / Chapter: {ref_chapter}"
                        )
                elif "社内学習" in task_type:
                    tool_used = normalized_row.get("使用ツール", "")
                    if tool_used:
                        appeal_info = f"[Appeal/Tool Used] {tool_used}"

                if appeal_info:
                    line_items.append(appeal_info)

                records.append(
                    NormalizedRecord(row_number=row_idx, data=normalized_row)
                )
                text_lines.append(f"[Row {row_idx}] " + ", ".join(line_items))

        plain_text = f"--- Sheet Name: {ws.title} ---\n" + "\n".join(text_lines)

        return SheetData(
            sheet_name=ws.title,
            is_visible=is_visible,
            headers=headers,
            records=records,
            plain_text_summary=plain_text,
        )

    def process_file(self, file_path: str) -> NormalizedWorkbook:
        path = Path(file_path)
        wb = openpyxl.load_workbook(filename=file_path, data_only=True)
        parsed_sheets: List[SheetData] = []

        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            sheet_data = self._extract_sheet(ws)
            if sheet_data:
                parsed_sheets.append(sheet_data)

        wb.close()

        return NormalizedWorkbook(
            file_name=path.name,
            processed_at=datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
            sheets=parsed_sheets,
        )