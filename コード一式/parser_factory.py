"""Parser factory module providing parser creation, management, and Excel normalization entry points."""

from abc import ABC, abstractmethod
import logging
from pathlib import Path
import sys
from typing import Any, Dict, List

import openpyxl
import pandas as pd

from excel_normalizer import ExcelNormalizer, NormalizedWorkbook

# Configure logging for VS Code output console
logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s [%(levelname)s] %(message)s",
    datefmt="%H:%M:%S",
)
logger = logging.getLogger(__name__)


# ==============================================================================
# 1. Abstract Base Class for Parsers
# ==============================================================================
class BaseExcelParser(ABC):
    """Abstract base class inherited by all data loading parsers."""

    @abstractmethod
    def parse(self, file_path: Path) -> List[Dict[str, Any]]:
        """Parse the given file and return a list of record dictionaries."""
        pass


# ==============================================================================
# 2. Concrete Parsers for Supported File Extensions
# ==============================================================================
class OpenPyXLParser(BaseExcelParser):
    """Parser supporting .xlsx format using openpyxl."""

    def __init__(self, max_empty_threshold: int = 3):
        """:param max_empty_threshold: Threshold of consecutive empty rows for early termination."""
        self.max_empty_threshold = max_empty_threshold

    def _unmerge_and_fill(self, ws) -> None:
        """Unmerge cells and propagate the top-left cell value to all cells in the range."""
        merged_ranges = list(ws.merged_cells.ranges)

        if merged_ranges:
            logger.info(
                f"   [Unmerge & Fill] Detected {len(merged_ranges)} merged ranges. Propagating values."
            )

        for cell_range in merged_ranges:
            top_left_val = ws.cell(
                row=cell_range.min_row, column=cell_range.min_col
            ).value

            ws.unmerge_cells(str(cell_range))

            for r in range(cell_range.min_row, cell_range.max_row + 1):
                for c in range(cell_range.min_col, cell_range.max_col + 1):
                    ws.cell(row=r, column=c).value = top_left_val

    def _get_cell_value(self, cell_data, cell_formula) -> Any:
        """Retrieve cell calculation result, logging formula cache misses safely."""
        val = cell_data.value

        if val is None and cell_formula.value is not None:
            formula_str = str(cell_formula.value)
            if formula_str.startswith("="):
                logger.warning(
                    f"   [Formula Cache Miss] Calculated value for cell {cell_data.coordinate} is missing "
                    f"(Formula: {formula_str}). Filling with empty string."
                )

        return val if val is not None else ""

    def parse(self, file_path: Path) -> List[Dict[str, Any]]:
        logger.info(f"[OpenPyXLParser] Starting processing: {file_path.name}")

        wb_data = openpyxl.load_workbook(file_path, data_only=True)
        wb_formula = openpyxl.load_workbook(file_path, data_only=False)

        ws_data = wb_data.active
        ws_formula = wb_formula.active

        self._unmerge_and_fill(ws_data)

        max_col = ws_data.max_column
        max_row = ws_data.max_row

        if max_row < 1 or max_col < 1:
            wb_data.close()
            wb_formula.close()
            return []

        header_map = {}
        for c_idx in range(1, max_col + 1):
            val = ws_data.cell(row=1, column=c_idx).value
            if val is not None and str(val).strip() != "":
                header_map[c_idx] = str(val).strip()

        if not header_map:
            for c_idx in range(1, max_col + 1):
                header_map[c_idx] = f"col_{c_idx}"

        parsed_data = []
        empty_row_count = 0
        has_started = False

        for r_idx in range(2, max_row + 1):
            row_all_values = [
                self._get_cell_value(
                    ws_data.cell(row=r_idx, column=c_idx),
                    ws_formula.cell(row=r_idx, column=c_idx),
                )
                for c_idx in range(1, max_col + 1)
            ]

            is_empty_row = all(v == "" or v is None for v in row_all_values)

            if is_empty_row:
                if not has_started:
                    continue

                empty_row_count += 1
                if empty_row_count >= self.max_empty_threshold:
                    logger.info(
                        f"   [Boundary Control] Reached {self.max_empty_threshold} consecutive empty rows. "
                        f"Terminating early at row {r_idx}."
                    )
                    break
                continue

            has_started = True
            empty_row_count = 0

            row_dict = {}
            for c_idx, header_name in header_map.items():
                c_data = ws_data.cell(row=r_idx, column=c_idx)
                c_form = ws_formula.cell(row=r_idx, column=c_idx)
                row_dict[header_name] = self._get_cell_value(c_data, c_form)

            parsed_data.append(row_dict)

        wb_data.close()
        wb_formula.close()
        return parsed_data


class XLRDParser(BaseExcelParser):
    """Parser supporting legacy .xls format using pandas and xlrd."""

    def parse(self, file_path: Path) -> List[Dict[str, Any]]:
        logger.info(f"[XLRDParser] Starting processing: {file_path.name}")
        df = pd.read_excel(file_path, engine="xlrd")
        df = df.ffill().fillna("")
        return df.to_dict(orient="records")


class CSVParser(BaseExcelParser):
    """Parser supporting .csv text format using pandas."""

    def parse(self, file_path: Path) -> List[Dict[str, Any]]:
        logger.info(f"[CSVParser] Starting processing: {file_path.name}")
        try:
            df = pd.read_csv(file_path, encoding="utf-8")
        except UnicodeDecodeError:
            logger.warning(
                "UTF-8 decoding failed. Retrying with CP932 (Shift-JIS)."
            )
            df = pd.read_csv(file_path, encoding="cp932")

        df = df.fillna("")
        return df.to_dict(orient="records")


# ==============================================================================
# 3. Factory Classes
# ==============================================================================
class ParserFactory:
    """Factory class for generating and managing parser instances."""

    @staticmethod
    def create_excel_normalizer(
        ignore_hidden_sheets: bool = True,
    ) -> ExcelNormalizer:
        """Create and return an instance of ExcelNormalizer."""
        return ExcelNormalizer(ignore_hidden_sheets=ignore_hidden_sheets)


class ExcelParserFactory:
    """Factory class to determine and instantiate the appropriate parser based on file extension."""

    _parsers = {
        ".xlsx": OpenPyXLParser,
        ".xls": XLRDParser,
        ".csv": CSVParser,
    }

    @classmethod
    def get_parser(cls, file_path: Path) -> BaseExcelParser:
        suffix = file_path.suffix.lower()
        parser_cls = cls._parsers.get(suffix)

        if not parser_cls:
            supported = ", ".join(cls._parsers.keys())
            raise ValueError(
                f"Unsupported file extension: '{suffix}' (Supported: {supported})"
            )

        return parser_cls()


# ==============================================================================
# 4. Main Entry Point for Testing & Execution
# ==============================================================================
if __name__ == "__main__":
    base_dir = Path(__file__).parent if "__file__" in locals() else Path.cwd()
    input_dir = base_dir / "inputs"

    logger.info("=== Parser Factory & Normalization Execution Test ===")

    if not input_dir.exists():
        input_dir.mkdir(parents=True, exist_ok=True)
        logger.info(f"Created test directory: {input_dir.resolve()}")
    else:
        target_files = [
            f
            for f in input_dir.glob("*")
            if f.suffix.lower() in [".xlsx", ".xls", ".csv"]
        ]

        if not target_files:
            logger.info(
                f"No supported files (.xlsx / .xls / .csv) found in '{input_dir}'."
            )
        else:
            for file_path in target_files:
                try:
                    parser = ExcelParserFactory.get_parser(file_path)
                    result = parser.parse(file_path)
                    logger.info(
                        f"Success: {file_path.name} -> Extracted {len(result)} rows."
                    )
                    if result:
                        logger.info(f"   [First Row Sample]: {result[0]}")
                except Exception as e:
                    logger.error(f"Failed: {file_path.name} -> Error: {e}")

    # Integrated execution entry point for workbook normalization
    target_file = base_dir / "フェーズ1_コアエンジン開発_ToDoリスト.xlsx"
    if target_file.exists():
        print(f"Processing target file via ParserFactory: {target_file.name}")
        try:
            normalizer = ParserFactory.create_excel_normalizer(
                ignore_hidden_sheets=True
            )
            workbook_ir: NormalizedWorkbook = normalizer.process_file(
                str(target_file)
            )

            print(workbook_ir.to_json())
            print("Successfully generated IR JSON via ParserFactory.")

        except Exception as e:
            print(f"Error occurred during execution: {e}", file=sys.stderr)
            sys.exit(1)